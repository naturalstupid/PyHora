"""Excel (.xlsx) reader for the batch layer.

Mirrors csv_reader: yields the same :class:`RawRow` objects so the rest of the
pipeline (BirthRecord.from_dict, engine) is reader-agnostic. The only extra work
here is turning native Excel cell types (numbers, dates, times) into the string
forms the shared validator expects — so a user can format the birth date as a
real date cell and the time as a real time cell instead of quoting text.
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

from .csv_reader import RawRow


def _stringify_cell(value) -> str:
    """Convert an Excel cell value to the canonical string BirthRecord expects.

    * date / date-only datetime -> ``YYYY,M,D``
    * time / time-only datetime (Excel's 1899 epoch) -> ``HH:MM:SS``
    * whole-number float -> integer text (27.0 -> "27")
    * other numbers -> full-precision text (75.783333 stays exact)
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, _dt.datetime):
        # Excel stores a time-only cell as a datetime on 1899-12-30/31.
        if value.year <= 1899:
            return value.strftime("%H:%M:%S")
        return f"{value.year},{value.month},{value.day}"
    if isinstance(value, _dt.date):
        return f"{value.year},{value.month},{value.day}"
    if isinstance(value, _dt.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def read_excel(path, *, sheet=None) -> Iterator[RawRow]:
    """Yield ``RawRow`` for each non-empty data row in an .xlsx worksheet.

    ``sheet`` selects a worksheet by name or index; default is the active sheet.
    The first non-empty row is the header; header names are lower-cased and
    stripped so the downstream alias map matches. Fully blank rows are skipped.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Excel file not found: {path}")

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        if sheet is None:
            ws = wb.active
        elif isinstance(sheet, int):
            ws = wb.worksheets[sheet]
        else:
            ws = wb[sheet]

        header = None
        row_number = 0
        for line_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [_stringify_cell(v) for v in row]
            if not any(v != "" for v in values):
                continue  # skip fully-blank row (also skips leading blanks)
            if header is None:
                header = [v.strip().lower() for v in values]
                continue
            # pad/truncate to header width
            if len(values) < len(header):
                values += [""] * (len(header) - len(values))
            data = {header[i]: values[i] for i in range(len(header)) if header[i]}
            row_number += 1
            yield RawRow(row_number=row_number, line_number=line_number, data=data)
    finally:
        wb.close()
